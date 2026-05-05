from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, time as dt_time, timedelta, timezone
from decimal import Decimal
from typing import Any, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.responses import StreamingResponse
from sqlalchemy import case, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.cashiers import build_cashier_shift_summary
from app.db.models import ActionLog, Cashier, CashierShift, Client, FuelSale, ReceiptCancellation, Transaction, TransactionType
from app.db.session import get_db
from app.services.admin_auth import require_admin
from app.services.cashier_auth import decode_cashier_token
from app.utils.datetime import excel_local_datetime, format_local_datetime, local_now, local_today


def _set_report_cache_headers(response: Response) -> None:
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"


router = APIRouter(prefix="/api/reports", tags=["reports"], dependencies=[Depends(_set_report_cache_headers)])
security = HTTPBearer(auto_error=False)

_MONTH_NAMES_RU = [
    "Январь",
    "Февраль",
    "Март",
    "Апрель",
    "Май",
    "Июнь",
    "Июль",
    "Август",
    "Сентябрь",
    "Октябрь",
    "Ноябрь",
    "Декабрь",
]


async def require_reports_access(
    request: Request,
    credentials: HTTPAuthorizationCredentials | None = Depends(security),
    db: AsyncSession = Depends(get_db),
) -> bool:
    try:
        require_admin(request=request, credentials=None)
        return True
    except HTTPException:
        pass

    if credentials is not None:
        cashier_id = decode_cashier_token(credentials.credentials)
        if cashier_id is not None:
            cashier = (await db.execute(select(Cashier).where(Cashier.id == cashier_id))).scalar_one_or_none()
            if cashier is not None and cashier.is_active:
                return True

    raise HTTPException(status_code=401, detail="Нет доступа к отчётам")


def _to_number(val: Any) -> float:
    if val is None:
        return 0.0
    if isinstance(val, Decimal):
        return float(val)
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(val)
    except Exception:
        return 0.0


def _excel_datetime(val: Any) -> Any:
    """openpyxl не поддерживает timezone-aware datetime."""
    if not isinstance(val, datetime):
        return val
    local_value = excel_local_datetime(val)
    return local_value if local_value is not None else val


def _dt_label(val: Optional[datetime]) -> str:
    return format_local_datetime(val)


def _report_datetime_bounds(date_from: date, date_to: date) -> tuple[datetime, datetime]:
    start = datetime.combine(date_from, dt_time.min)
    end = datetime.combine(date_to + timedelta(days=1), dt_time.min)
    return start, end


def _cancelled_fuel_sale_clause():
    return select(ReceiptCancellation.id).where(
        or_(
            FuelSale.check_id == ReceiptCancellation.original_check_id,
            FuelSale.check_id.like(ReceiptCancellation.original_check_id.concat("-C-F%")),
        )
    ).exists()


def _cancelled_transaction_clause():
    return select(ReceiptCancellation.id).where(
        or_(
            Transaction.check_id == ReceiptCancellation.original_check_id,
            Transaction.check_id == ReceiptCancellation.original_check_id.concat("-R"),
            Transaction.check_id == ReceiptCancellation.original_check_id.concat("-C-A"),
            Transaction.check_id == ReceiptCancellation.original_check_id.concat("-C-R"),
        )
    ).exists()


def _resolve_admin_summary_period(
    period_kind: str,
    year: Optional[int],
    month: Optional[int],
) -> tuple[str, int, Optional[int], date, date, str]:
    today = local_today()
    normalized_kind = (period_kind or "month").strip().lower()
    if normalized_kind not in {"month", "year"}:
        raise HTTPException(status_code=400, detail="Период должен быть month или year")

    resolved_year = year or today.year
    if normalized_kind == "year":
        date_from = date(resolved_year, 1, 1)
        date_to = date(resolved_year, 12, 31)
        return normalized_kind, resolved_year, None, date_from, date_to, f"{resolved_year} год"

    resolved_month = month or today.month
    try:
        date_from = date(resolved_year, resolved_month, 1)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Некорректный месяц или год") from exc

    if resolved_month == 12:
        next_month = date(resolved_year + 1, 1, 1)
    else:
        next_month = date(resolved_year, resolved_month + 1, 1)
    date_to = next_month - timedelta(days=1)
    period_label = f"{_MONTH_NAMES_RU[resolved_month - 1]} {resolved_year}"
    return normalized_kind, resolved_year, resolved_month, date_from, date_to, period_label


@router.get("/cashiers")
async def report_cashiers(
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Cashier.id, Cashier.name)
        .where(Cashier.is_active.is_(True))
        .order_by(Cashier.name.asc(), Cashier.id.asc())
    )
    return [{"id": row.id, "name": row.name} for row in result.all()]


@router.get("/admin-summary")
async def report_admin_summary(
    period_kind: str = Query("month"),
    year: Optional[int] = Query(None, ge=2000, le=2100),
    month: Optional[int] = Query(None, ge=1, le=12),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    resolved_kind, resolved_year, resolved_month, date_from, date_to, period_label = _resolve_admin_summary_period(
        period_kind,
        year,
        month,
    )
    dt_from, dt_to_exclusive = _report_datetime_bounds(date_from, date_to)

    fuel_result = await db.execute(
        select(
            FuelSale.fuel_type,
            func.sum(FuelSale.liters).label("total_liters"),
        )
        .where(
            FuelSale.sale_date >= dt_from,
            FuelSale.sale_date < dt_to_exclusive,
            ~_cancelled_fuel_sale_clause(),
        )
        .group_by(FuelSale.fuel_type)
        .order_by(func.sum(FuelSale.liters).desc(), FuelSale.fuel_type.asc())
    )
    fuel_rows = [
        {
            "fuel_type": row.fuel_type or "Неизвестно",
            "total_liters": _to_number(row.total_liters),
        }
        for row in fuel_result.all()
    ]

    bonus_result = await db.execute(
        select(
            func.sum(
                case(
                    (Transaction.type == TransactionType.accrual, Transaction.amount_bonus),
                    else_=0,
                )
            ).label("accrued_bonus_total"),
            func.sum(
                case(
                    (Transaction.type == TransactionType.redemption, Transaction.amount_bonus),
                    else_=0,
                )
            ).label("redeemed_bonus_total"),
        )
        .where(
            Transaction.ts >= dt_from,
            Transaction.ts < dt_to_exclusive,
            ~_cancelled_transaction_clause(),
        )
    )
    bonus_row = bonus_result.one()

    totals = {
        "fuel_types_count": len(fuel_rows),
        "total_liters": sum(row["total_liters"] for row in fuel_rows),
        "accrued_bonus_total": _to_number(bonus_row.accrued_bonus_total),
        "redeemed_bonus_total": _to_number(bonus_row.redeemed_bonus_total),
    }

    payload = {
        "period_kind": resolved_kind,
        "period_label": period_label,
        "year": resolved_year,
        "month": resolved_month,
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "fuel_rows": fuel_rows,
        "totals": totals,
    }

    export_rows = [
        {
            "Период": period_label,
            "Топливо": row["fuel_type"],
            "Литров": row["total_liters"],
            "Начислено бонусов": "",
            "Списано бонусов": "",
        }
        for row in fuel_rows
    ]
    export_rows.append(
        {
            "Период": period_label,
            "Топливо": "ИТОГО",
            "Литров": totals["total_liters"],
            "Начислено бонусов": totals["accrued_bonus_total"],
            "Списано бонусов": totals["redeemed_bonus_total"],
        }
    )

    if fmt == "csv":
        return _stream_csv(export_rows, "admin_summary.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows,
            filename="admin_summary.xlsx",
            sheet_name="Краткий отчет",
            columns=["Период", "Топливо", "Литров", "Начислено бонусов", "Списано бонусов"],
            number_formats={
                "Литров": "#,##0.000",
                "Начислено бонусов": "#,##0.00",
                "Списано бонусов": "#,##0.00",
            },
        )
    return payload


_ACTION_LOG_ACTOR_LABELS = {"admin": "Админ", "cashier": "Кассир", "system": "Система"}
_ACTION_LOG_ACTION_LABELS = {
    "admin_login": "Вход админа",
    "admin_logout": "Выход админа",
    "cashier_login": "Вход кассира",
    "cashier_logout": "Выход кассира",
    "shift_opened": "Открытие смены",
    "shift_closed": "Закрытие смены",
    "receipt_accrued": "Начисление по чеку",
    "receipt_redeemed": "Списание по чеку",
    "receipt_cancelled": "Отмена чека",
    "fuel_sale_logged": "Запись продажи топлива",
    "client_registered_cashier": "Регистрация клиента с кассы",
    "news_created": "Создание новости",
    "news_updated": "Изменение новости",
    "news_status_updated": "Изменение статуса новости",
    "news_deleted": "Удаление новости",
    "rule_created": "Создание правила",
    "rule_updated": "Изменение правила",
    "rule_deleted": "Удаление правила",
    "campaign_created": "Создание акции",
    "campaign_updated": "Изменение акции",
    "campaign_deleted": "Удаление акции",
    "cashier_created": "Создание кассира",
    "cashier_updated": "Изменение кассира",
    "cashier_archived": "Архивация кассира",
    "cashier_deleted": "Удаление кассира",
    "client_password_updated": "Изменение пароля клиента",
}
_ACTION_LOG_ENTITY_LABELS = {
    "session": "Сессия",
    "cashier_shift": "Смена",
    "transaction": "Операция",
    "receipt": "Чек",
    "fuel_sale": "Продажа топлива",
    "client": "Клиент",
    "news": "Новость",
    "rule": "Правило",
    "campaign": "Акция",
    "cashier": "Кассир",
}
_ACTION_LOG_LOCATION_LABELS = {"base": "Магазин", "fuel": "Топливо"}
_ACTION_LOG_TYPE_LABELS = {"percent": "Процент", "fixed": "Фиксированное значение"}
_ACTION_LOG_DETAIL_LABELS = {
    "title": "Заголовок",
    "name": "Название",
    "phone": "Телефон",
    "purchase_amount": "Сумма чека",
    "amount_bonus": "Бонусы",
    "location": "Место",
    "fuel_type": "Топливо",
    "liters": "Литры",
    "total_rub": "Сумма",
    "original_check_id": "Исходный чек",
    "canceled_check_id": "Чек отмены",
    "restored_bonus": "Возвращено бонусов",
    "deducted_bonus": "Снято бонусов",
    "reversed_transactions": "Обратных бонусных операций",
    "reversed_fuel_sales": "Обратных топливных продаж",
    "reason": "Причина",
    "type": "Тип начисления",
    "username": "Логин",
    "is_active": "Активен",
    "transactions_count": "Операций",
    "accrual_transactions_count": "Начислений",
    "redemption_transactions_count": "Списаний",
    "accrued_bonus_total": "Начислено бонусов",
    "redeemed_bonus_total": "Списано бонусов",
    "fuel_sales_count": "Продаж топлива",
    "fuel_sales_total": "Сумма топлива",
    "fuel_liters_total": "Литров топлива",
}


def _format_number_text(value: Any, digits: int = 2) -> str:
    return f"{_to_number(value):,.{digits}f}".replace(",", " ")


def _format_integer_text(value: Any) -> str:
    return f"{int(round(_to_number(value))):,}".replace(",", " ")


def _action_log_actor_display(actor_type: str, actor_name: str) -> str:
    actor_label = _ACTION_LOG_ACTOR_LABELS.get(actor_type, actor_type)
    actor_name = (actor_name or "").strip()
    if actor_name:
        return f"{actor_label}: {actor_name}"
    return actor_label


def _action_log_parse_details(details: Optional[str]) -> Any:
    if not details:
        return None
    try:
        return json.loads(details)
    except Exception:
        return details


def _action_log_format_value(key: str, value: Any) -> Optional[str]:
    if value in (None, "", [], {}):
        return None
    if key in {
        "purchase_amount",
        "amount_bonus",
        "total_rub",
        "restored_bonus",
        "deducted_bonus",
        "accrued_bonus_total",
        "redeemed_bonus_total",
        "fuel_sales_total",
    }:
        return _format_number_text(value, 2)
    if key in {"liters", "fuel_liters_total"}:
        return _format_number_text(value, 3)
    if key in {
        "reversed_transactions",
        "reversed_fuel_sales",
        "transactions_count",
        "accrual_transactions_count",
        "redemption_transactions_count",
        "fuel_sales_count",
    }:
        return _format_integer_text(value)
    if key == "location":
        return _ACTION_LOG_LOCATION_LABELS.get(str(value), str(value))
    if key == "type":
        return _ACTION_LOG_TYPE_LABELS.get(str(value), str(value))
    if key == "is_active":
        return "Да" if bool(value) else "Нет"
    return str(value)


def _action_log_entity_display(
    *,
    action: str,
    entity_type: Optional[str],
    entity_id: Optional[int],
    check_id: Optional[str],
    details: Any,
) -> str:
    if check_id:
        if action == "fuel_sale_logged":
            return f"Продажа топлива по чеку {check_id}"
        return f"Чек {check_id}"

    details_dict = details if isinstance(details, dict) else {}
    label = _ACTION_LOG_ENTITY_LABELS.get(entity_type or "", entity_type or "Объект")

    if entity_type == "news" and details_dict.get("title"):
        return f"Новость: {details_dict['title']}"
    if entity_type == "campaign" and details_dict.get("name"):
        return f"Акция: {details_dict['name']}"
    if entity_type == "cashier" and details_dict.get("name"):
        return f"Кассир: {details_dict['name']}"
    if entity_type == "client":
        client_name = (details_dict.get("name") or "").strip()
        client_phone = (details_dict.get("phone") or "").strip()
        if client_name and client_phone:
            return f"Клиент: {client_name} ({client_phone})"
        if client_phone:
            return f"Клиент: {client_phone}"
        if client_name:
            return f"Клиент: {client_name}"

    if entity_id is not None:
        return f"{label} #{entity_id}"
    return label


def _action_log_details_text(action: str, details: Any) -> str:
    if details in (None, "", [], {}):
        return ""
    if isinstance(details, str):
        return details
    if not isinstance(details, dict):
        return json.dumps(details, ensure_ascii=False, default=str)

    ordered_fields: dict[str, list[tuple[str, str]]] = {
        "receipt_accrued": [
            ("purchase_amount", "Сумма чека"),
            ("amount_bonus", "Начислено бонусов"),
            ("location", "Место"),
        ],
        "receipt_redeemed": [
            ("purchase_amount", "Сумма чека"),
            ("amount_bonus", "Списано бонусов"),
            ("location", "Место"),
        ],
        "fuel_sale_logged": [
            ("fuel_type", "Топливо"),
            ("liters", "Литры"),
            ("total_rub", "Сумма"),
        ],
        "client_registered_cashier": [
            ("phone", "Телефон"),
            ("name", "Имя"),
        ],
        "receipt_cancelled": [
            ("canceled_check_id", "Чек отмены"),
            ("purchase_amount", "Сумма чека"),
            ("deducted_bonus", "Снято бонусов"),
            ("restored_bonus", "Возвращено бонусов"),
            ("reversed_transactions", "Обратных бонусных операций"),
            ("reversed_fuel_sales", "Обратных топливных продаж"),
            ("reason", "Причина"),
        ],
        "shift_closed": [
            ("transactions_count", "Операций"),
            ("accrual_transactions_count", "Начислений"),
            ("redemption_transactions_count", "Списаний"),
            ("accrued_bonus_total", "Начислено бонусов"),
            ("redeemed_bonus_total", "Списано бонусов"),
            ("fuel_sales_count", "Продаж топлива"),
            ("fuel_sales_total", "Сумма топлива"),
            ("fuel_liters_total", "Литров топлива"),
        ],
        "cashier_logout": [
            ("transactions_count", "Операций"),
            ("accrued_bonus_total", "Начислено бонусов"),
            ("redeemed_bonus_total", "Списано бонусов"),
            ("fuel_sales_count", "Продаж топлива"),
        ],
        "news_created": [("title", "Заголовок")],
        "news_updated": [("title", "Заголовок"), ("is_active", "Активна")],
        "news_status_updated": [("title", "Заголовок"), ("is_active", "Активна")],
        "news_deleted": [("title", "Заголовок")],
        "rule_created": [("location", "Место"), ("type", "Тип начисления")],
        "rule_updated": [("location", "Место"), ("type", "Тип начисления")],
        "rule_deleted": [("location", "Место"), ("type", "Тип начисления")],
        "campaign_created": [("name", "Название")],
        "campaign_updated": [("name", "Название")],
        "campaign_deleted": [("name", "Название")],
        "cashier_created": [("name", "Имя"), ("username", "Логин")],
        "cashier_updated": [("name", "Имя"), ("username", "Логин"), ("is_active", "Активен")],
        "cashier_archived": [("name", "Имя")],
        "cashier_deleted": [("name", "Имя")],
    }

    lines: list[str] = []
    used_keys: set[str] = set()
    for key, label in ordered_fields.get(action, []):
        formatted_value = _action_log_format_value(key, details.get(key))
        if formatted_value is None:
            continue
        lines.append(f"{label}: {formatted_value}")
        used_keys.add(key)

    for key, value in details.items():
        if key in used_keys:
            continue
        formatted_value = _action_log_format_value(key, value)
        if formatted_value is None:
            continue
        label = _ACTION_LOG_DETAIL_LABELS.get(key, key.replace("_", " "))
        lines.append(f"{label}: {formatted_value}")

    return "\n".join(lines)


def _stream_csv(rows: list[dict], filename: str) -> StreamingResponse:
    if not rows:
        content = "\ufeff"  # BOM only
    else:
        buf = io.StringIO()
        buf.write("\ufeff")  # UTF-8 BOM for Excel
        # В RU-локали Excel чаще ожидает ';' как разделитель
        writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
        content = buf.getvalue()
    return StreamingResponse(
        iter([content.encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _stream_xlsx(
    *,
    rows: list[dict],
    filename: str,
    sheet_name: str,
    columns: list[str],
    number_formats: Optional[dict[str, str]] = None,
) -> StreamingResponse:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=(
                "Excel экспорт недоступен: не установлен пакет openpyxl. "
                "Установите зависимости проекта и перезапустите сервер."
            ),
        ) from e

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Report")[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    # Header
    for col_idx, key in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=key)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    ws.freeze_panes = "A2"

    # Data
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(columns, start=1):
            val = row.get(key, "")

            # Нормализация типов для openpyxl/Excel
            tz = getattr(val, "tzinfo", None)
            if tz is not None:
                if isinstance(val, datetime):
                    val = val.astimezone(timezone.utc).replace(tzinfo=None)
                elif isinstance(val, dt_time):
                    val = val.replace(tzinfo=None)
                else:
                    # На случай нестандартных datetime-подобных типов
                    try:
                        val = val.replace(tzinfo=None)
                    except Exception:
                        pass
            if isinstance(val, Decimal):
                val = float(val)

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            fmt = (number_formats or {}).get(key)
            if fmt:
                cell.number_format = fmt
            if isinstance(val, (int, float, Decimal)):
                cell.alignment = right
            else:
                cell.alignment = left

    # Autofilter
    last_row = max(1, 1 + len(rows))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"

    # Column widths (simple autosize)
    for col_idx, key in enumerate(columns, start=1):
        max_len = len(str(key))
        for row in rows[:5000]:
            v = row.get(key, "")
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(42, max_len + 2))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/cashier-sales")
async def report_cashier_sales(
    date_from: date = Query(...),
    date_to: date = Query(...),
    cashier_id: Optional[int] = Query(None),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    dt_from, dt_to_exclusive = _report_datetime_bounds(date_from, date_to)
    fuel_stmt = (
        select(
            FuelSale.cashier_id.label("cashier_id"),
            func.coalesce(Cashier.name, FuelSale.cashier_name).label("cashier_name"),
            func.count(FuelSale.id).label("sales_count"),
            func.sum(FuelSale.total_rub).label("total_rub"),
            func.sum(FuelSale.liters).label("total_liters"),
        )
        .select_from(FuelSale)
        .outerjoin(Cashier, Cashier.id == FuelSale.cashier_id)
        .where(
            FuelSale.sale_date >= dt_from,
            FuelSale.sale_date < dt_to_exclusive,
            ~_cancelled_fuel_sale_clause(),
        )
        .group_by(FuelSale.cashier_id, Cashier.name, FuelSale.cashier_name)
    )
    if cashier_id is not None:
        fuel_stmt = fuel_stmt.where(FuelSale.cashier_id == cashier_id)

    tx_stmt = (
        select(
            Transaction.cashier_id.label("cashier_id"),
            Cashier.name.label("cashier_name"),
            func.count(Transaction.id).label("transactions_count"),
            func.sum(
                case(
                    (Transaction.type == TransactionType.accrual, Transaction.amount_bonus),
                    else_=0,
                )
            ).label("accrued_bonus_total"),
            func.sum(
                case(
                    (Transaction.type == TransactionType.redemption, Transaction.amount_bonus),
                    else_=0,
                )
            ).label("redeemed_bonus_total"),
            func.sum(
                case(
                    (Transaction.type == TransactionType.accrual, 1),
                    else_=0,
                )
            ).label("accrual_count"),
            func.sum(
                case(
                    (Transaction.type == TransactionType.redemption, 1),
                    else_=0,
                )
            ).label("redemption_count"),
        )
        .select_from(Transaction)
        .outerjoin(Cashier, Cashier.id == Transaction.cashier_id)
        .where(
            Transaction.ts >= dt_from,
            Transaction.ts < dt_to_exclusive,
            Transaction.cashier_id.is_not(None),
            ~_cancelled_transaction_clause(),
        )
        .group_by(Transaction.cashier_id, Cashier.name)
    )
    if cashier_id is not None:
        tx_stmt = tx_stmt.where(Transaction.cashier_id == cashier_id)

    fuel_rows = await db.execute(fuel_stmt)
    tx_rows = await db.execute(tx_stmt)

    merged: dict[str, dict[str, Any]] = {}

    for r in fuel_rows.all():
        key = str(r.cashier_id) if r.cashier_id is not None else f"name:{r.cashier_name or '—'}"
        merged[key] = {
            "cashier_id": r.cashier_id,
            "cashier_name": r.cashier_name or "—",
            "sales_count": int(r.sales_count or 0),
            "total_rub": _to_number(r.total_rub),
            "total_liters": _to_number(r.total_liters),
            "transactions_count": 0,
            "accrual_count": 0,
            "redemption_count": 0,
            "accrued_bonus_total": 0.0,
            "redeemed_bonus_total": 0.0,
            "bonus_net_total": 0.0,
        }

    for r in tx_rows.all():
        key = str(r.cashier_id)
        row = merged.get(key)
        if row is None:
            row = {
                "cashier_id": r.cashier_id,
                "cashier_name": r.cashier_name or f"Кассир #{r.cashier_id}",
                "sales_count": 0,
                "total_rub": 0.0,
                "total_liters": 0.0,
                "transactions_count": 0,
                "accrual_count": 0,
                "redemption_count": 0,
                "accrued_bonus_total": 0.0,
                "redeemed_bonus_total": 0.0,
                "bonus_net_total": 0.0,
            }
            merged[key] = row
        row["transactions_count"] = int(r.transactions_count or 0)
        row["accrual_count"] = int(r.accrual_count or 0)
        row["redemption_count"] = int(r.redemption_count or 0)
        row["accrued_bonus_total"] = _to_number(r.accrued_bonus_total)
        row["redeemed_bonus_total"] = _to_number(r.redeemed_bonus_total)
        row["bonus_net_total"] = row["accrued_bonus_total"] - row["redeemed_bonus_total"]

    json_rows = sorted(merged.values(), key=lambda item: ((item["cashier_name"] or "").lower(), item["cashier_id"] or 0))

    export_rows = [
        {
            "Кассир": r["cashier_name"],
            "Операций с бонусами": r["transactions_count"],
            "Начислений": r["accrual_count"],
            "Списаний": r["redemption_count"],
            "Начислено бонусов": r["accrued_bonus_total"],
            "Списано бонусов": r["redeemed_bonus_total"],
            "Итог бонусов": r["bonus_net_total"],
            "Продаж": r["sales_count"],
            "Сумма, ₽": r["total_rub"],
            "Литров": r["total_liters"],
        }
        for r in json_rows
    ]

    if export_rows:
        export_rows.append(
            {
                "Кассир": "ИТОГО",
                "Операций с бонусами": sum(int(r["transactions_count"]) for r in json_rows),
                "Начислений": sum(int(r["accrual_count"]) for r in json_rows),
                "Списаний": sum(int(r["redemption_count"]) for r in json_rows),
                "Начислено бонусов": sum(float(r["accrued_bonus_total"]) for r in json_rows),
                "Списано бонусов": sum(float(r["redeemed_bonus_total"]) for r in json_rows),
                "Итог бонусов": sum(float(r["bonus_net_total"]) for r in json_rows),
                "Продаж": sum(int(r["sales_count"]) for r in json_rows),
                "Сумма, ₽": sum(float(r["total_rub"]) for r in json_rows),
                "Литров": sum(float(r["total_liters"]) for r in json_rows),
            }
        )

    if fmt == "csv":
        return _stream_csv(export_rows, "cashier_sales.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows,
            filename="cashier_sales.xlsx",
            sheet_name="Продажи кассиров",
            columns=[
                "Кассир",
                "Операций с бонусами",
                "Начислений",
                "Списаний",
                "Начислено бонусов",
                "Списано бонусов",
                "Итог бонусов",
                "Продаж",
                "Сумма, ₽",
                "Литров",
            ],
            number_formats={
                "Операций с бонусами": "0",
                "Начислений": "0",
                "Списаний": "0",
                "Начислено бонусов": "#,##0.00",
                "Списано бонусов": "#,##0.00",
                "Итог бонусов": "#,##0.00",
                "Продаж": "0",
                "Сумма, ₽": "#,##0.00",
                "Литров": "#,##0.000",
            },
        )
    return json_rows


@router.get("/cashier-hours")
async def report_cashier_hours(
    date_from: date = Query(...),
    date_to: date = Query(...),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    """Отработанные часы кассиров по сменам (login/logout)."""
    dt_from, dt_to_exclusive = _report_datetime_bounds(date_from, date_to)
    now = local_now()

    # Смена попадает в отчёт если пересекает интервал
    effective_end = func.least(func.coalesce(CashierShift.ended_at, now), dt_to_exclusive)
    effective_start = func.greatest(CashierShift.started_at, dt_from)
    seconds_expr = func.extract("epoch", effective_end - effective_start)

    stmt = (
        select(
            Cashier.id.label("cashier_id"),
            Cashier.name.label("cashier_name"),
            func.count(CashierShift.id).label("shifts_count"),
            func.sum(func.greatest(seconds_expr, 0)).label("seconds_total"),
        )
        .select_from(CashierShift)
        .join(Cashier, Cashier.id == CashierShift.cashier_id)
        .where(
            CashierShift.started_at < dt_to_exclusive,
            func.coalesce(CashierShift.ended_at, now) >= dt_from,
        )
        .group_by(Cashier.id, Cashier.name)
        .order_by(Cashier.name)
    )
    result = await db.execute(stmt)
    rows = []
    for r in result.all():
        seconds_total = float(r.seconds_total or 0)
        hours_total = round(seconds_total / 3600.0, 2)
        rows.append(
            {
                "cashier_id": r.cashier_id,
                "cashier_name": r.cashier_name,
                "shifts_count": int(r.shifts_count or 0),
                "hours_total": hours_total,
            }
        )

    export_rows = [
        {
            "Кассир": r["cashier_name"],
            "Смен": r["shifts_count"],
            "Часов": r["hours_total"],
        }
        for r in rows
    ]

    if fmt == "csv":
        return _stream_csv(export_rows, "cashier_hours.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows,
            filename="cashier_hours.xlsx",
            sheet_name="Часы кассиров",
            columns=["Кассир", "Смен", "Часов"],
            number_formats={
                "Смен": "0",
                "Часов": "0.00",
            },
        )
    return rows


@router.get("/closed-shifts")
async def report_closed_shifts(
    date_from: date = Query(...),
    date_to: date = Query(...),
    cashier_id: Optional[int] = Query(None),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    dt_from, dt_to_exclusive = _report_datetime_bounds(date_from, date_to)
    stmt = (
        select(CashierShift, Cashier.name)
        .join(Cashier, Cashier.id == CashierShift.cashier_id)
        .where(
            CashierShift.ended_at.is_not(None),
            CashierShift.ended_at >= dt_from,
            CashierShift.ended_at < dt_to_exclusive,
        )
        .order_by(CashierShift.ended_at.desc(), CashierShift.id.desc())
    )
    if cashier_id is not None:
        stmt = stmt.where(CashierShift.cashier_id == cashier_id)
    result = await db.execute(stmt)
    shift_rows = result.all()
    summary_by_shift_id: dict[int, dict[str, Any]] = {}
    shift_ids = [shift.id for shift, _ in shift_rows]
    if shift_ids:
        logs_result = await db.execute(
            select(ActionLog.entity_id, ActionLog.details)
            .where(
                ActionLog.action == "shift_closed",
                ActionLog.entity_type == "cashier_shift",
                ActionLog.entity_id.in_(shift_ids),
            )
            .order_by(ActionLog.id.desc())
        )
        for entity_id, details in logs_result.all():
            if entity_id in summary_by_shift_id:
                continue
            if not details:
                continue
            try:
                summary_by_shift_id[int(entity_id)] = json.loads(details)
            except Exception:
                continue

    rows = []
    for shift, cashier_name in shift_rows:
        summary = summary_by_shift_id.get(shift.id) or await build_cashier_shift_summary(db, shift)
        duration_hours = 0.0
        if shift.started_at and shift.ended_at:
            duration_hours = round(max(0.0, (shift.ended_at - shift.started_at).total_seconds()) / 3600.0, 2)
        rows.append(
            {
                "shift_id": shift.id,
                "cashier_name": cashier_name,
                "started_at": _dt_label(shift.started_at),
                "ended_at": _dt_label(shift.ended_at),
                "duration_hours": duration_hours,
                "transactions_count": int(summary["transactions_count"]),
                "accrued_bonus_total": float(summary["accrued_bonus_total"]),
                "redeemed_bonus_total": float(summary["redeemed_bonus_total"]),
                "fuel_sales_count": int(summary["fuel_sales_count"]),
                "fuel_sales_total": float(summary["fuel_sales_total"]),
                "fuel_liters_total": float(summary["fuel_liters_total"]),
            }
        )

    export_rows = [
        {
            "Смена": row["shift_id"],
            "Кассир": row["cashier_name"],
            "Открыта": row["started_at"],
            "Закрыта": row["ended_at"],
            "Часов": row["duration_hours"],
            "Операций": row["transactions_count"],
            "Начислено бонусов": row["accrued_bonus_total"],
            "Списано бонусов": row["redeemed_bonus_total"],
            "Продаж топлива": row["fuel_sales_count"],
            "Топливо, сумма": row["fuel_sales_total"],
            "Топливо, литров": row["fuel_liters_total"],
        }
        for row in rows
    ]

    if fmt == "csv":
        return _stream_csv(export_rows, "closed_shifts.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows,
            filename="closed_shifts.xlsx",
            sheet_name="Закрытые смены",
            columns=[
                "Смена",
                "Кассир",
                "Открыта",
                "Закрыта",
                "Часов",
                "Операций",
                "Начислено бонусов",
                "Списано бонусов",
                "Продаж топлива",
                "Топливо, сумма",
                "Топливо, литров",
            ],
            number_formats={
                "Смена": "0",
                "Часов": "0.00",
                "Операций": "0",
                "Начислено бонусов": "#,##0.00",
                "Списано бонусов": "#,##0.00",
                "Продаж топлива": "0",
                "Топливо, сумма": "#,##0.00",
                "Топливо, литров": "#,##0.000",
            },
        )
    return rows


@router.get("/action-log")
async def report_action_log(
    date_from: date = Query(...),
    date_to: date = Query(...),
    limit: Optional[int] = Query(None, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    dt_from = datetime.combine(date_from, dt_time.min)
    dt_to_exclusive = datetime.combine(date_to + timedelta(days=1), dt_time.min)
    effective_limit = limit if limit is not None else (1000 if fmt in {"csv", "xlsx"} else 200)
    stmt = (
        select(
            ActionLog.created_at,
            ActionLog.actor_type,
            ActionLog.actor_name,
            ActionLog.action,
            ActionLog.entity_type,
            ActionLog.entity_id,
            ActionLog.check_id,
            ActionLog.details,
        )
        .where(ActionLog.created_at >= dt_from, ActionLog.created_at < dt_to_exclusive)
        .order_by(ActionLog.created_at.desc(), ActionLog.id.desc())
        .offset(offset)
        .limit(effective_limit)
    )
    result = await db.execute(stmt)
    rows = []
    for item in result.all():
        parsed_details = _action_log_parse_details(item.details)
        actor_type_label = _ACTION_LOG_ACTOR_LABELS.get(item.actor_type, item.actor_type)
        action_label = _ACTION_LOG_ACTION_LABELS.get(item.action, item.action)
        entity_type_label = _ACTION_LOG_ENTITY_LABELS.get(item.entity_type or "", item.entity_type or "")
        rows.append(
            {
                "created_at": _dt_label(item.created_at),
                "actor_type": actor_type_label,
                "actor_name": item.actor_name,
                "actor_display": _action_log_actor_display(item.actor_type, item.actor_name),
                "action": action_label,
                "entity_type": entity_type_label,
                "entity_id": item.entity_id,
                "entity_display": _action_log_entity_display(
                    action=item.action,
                    entity_type=item.entity_type,
                    entity_id=item.entity_id,
                    check_id=item.check_id,
                    details=parsed_details,
                ),
                "check_id": item.check_id or "",
                "details": _action_log_details_text(item.action, parsed_details),
            }
        )

    export_rows = [
        {
            "Дата": row["created_at"],
            "Кто": row["actor_display"],
            "Действие": row["action"],
            "Объект": row["entity_display"],
            "Чек": row["check_id"],
            "Подробности": row["details"],
        }
        for row in rows
    ]

    if fmt == "csv":
        return _stream_csv(export_rows, "action_log.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows,
            filename="action_log.xlsx",
            sheet_name="Журнал действий",
            columns=["Дата", "Кто", "Действие", "Объект", "Чек", "Подробности"],
        )
    return rows


@router.get("/fuel-sales")
async def report_fuel_sales(
    date_from: date = Query(...),
    date_to: date = Query(...),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    from sqlalchemy.orm import joinedload

    dt_from, dt_to_exclusive = _report_datetime_bounds(date_from, date_to)
    stmt = (
        select(FuelSale)
        .options(joinedload(FuelSale.client))
        .where(
            FuelSale.sale_date >= dt_from,
            FuelSale.sale_date < dt_to_exclusive,
            ~_cancelled_fuel_sale_clause(),
        )
        .order_by(FuelSale.sale_date.desc())
        .limit(500)
    )
    result = await db.execute(stmt)

    json_rows = []
    export_rows_csv = []
    export_rows_xlsx = []
    for sale in result.scalars().all():
        sale_date_str = _dt_label(sale.sale_date)
        client_phone = sale.client.phone if sale.client else ""
        client_name = sale.client.name if sale.client and sale.client.name else "Без клиента"
        liters = _to_number(sale.liters)
        price_per_liter = _to_number(sale.price_per_liter)
        total_rub = _to_number(sale.total_rub)

        json_rows.append(
            {
                "sale_date": sale_date_str,
                "cashier_name": sale.cashier_name or "—",
                "fuel_type": sale.fuel_type or "Неизвестно",
                "liters": liters,
                "price_per_liter": price_per_liter,
                "total_rub": total_rub,
                "client_phone": client_phone,
                "client_name": client_name,
            }
        )
        export_rows_csv.append(
            {
                "Дата": sale_date_str,
                "Кассир": sale.cashier_name or "—",
                "Топливо": sale.fuel_type or "Неизвестно",
                "Литров": liters,
                "Цена за литр, ₽": price_per_liter,
                "Сумма, ₽": total_rub,
                "Телефон": client_phone,
                "Клиент": client_name,
            }
        )
        export_rows_xlsx.append(
            {
                "Дата": _excel_datetime(sale.sale_date) if sale.sale_date else sale_date_str,
                "Кассир": sale.cashier_name or "—",
                "Топливо": sale.fuel_type or "Неизвестно",
                "Литров": liters,
                "Цена за литр, ₽": price_per_liter,
                "Сумма, ₽": total_rub,
                "Телефон": client_phone,
                "Клиент": client_name,
            }
        )

    if fmt == "csv":
        return _stream_csv(export_rows_csv, "fuel_sales.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows_xlsx,
            filename="fuel_sales.xlsx",
            sheet_name="Продажи топлива",
            columns=["Дата", "Кассир", "Топливо", "Литров", "Цена за литр, ₽", "Сумма, ₽", "Телефон", "Клиент"],
            number_formats={
                "Дата": "dd.mm.yyyy hh:mm",
                "Литров": "#,##0.000",
                "Цена за литр, ₽": "#,##0.00",
                "Сумма, ₽": "#,##0.00",
            },
        )
    return json_rows


@router.get("/fuel-by-type")
async def report_fuel_by_type(
    date_from: date = Query(...),
    date_to: date = Query(...),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    dt_from, dt_to_exclusive = _report_datetime_bounds(date_from, date_to)
    stmt = (
        select(
            FuelSale.fuel_type,
            func.count(FuelSale.id).label("sales_count"),
            func.sum(FuelSale.liters).label("total_liters"),
            func.sum(FuelSale.total_rub).label("total_rub"),
            func.avg(FuelSale.price_per_liter).label("avg_price_per_liter"),
            func.max(FuelSale.sale_date).label("last_sale_at"),
        )
        .where(
            FuelSale.sale_date >= dt_from,
            FuelSale.sale_date < dt_to_exclusive,
            ~_cancelled_fuel_sale_clause(),
        )
        .group_by(FuelSale.fuel_type)
        .order_by(func.max(FuelSale.sale_date).desc(), FuelSale.fuel_type.asc())
    )
    result = await db.execute(stmt)
    json_rows = [
        {
            "fuel_type": r.fuel_type or "Неизвестно",
            "sales_count": int(r.sales_count or 0),
            "total_liters": _to_number(r.total_liters),
            "total_rub": _to_number(r.total_rub),
            "avg_price_per_liter": _to_number(r.avg_price_per_liter),
            "last_sale_at": _dt_label(r.last_sale_at),
        }
        for r in result.all()
    ]

    export_rows = [
        {
            "Вид топлива": r["fuel_type"],
            "Продаж": r["sales_count"],
            "Литров": r["total_liters"],
            "Сумма, ₽": r["total_rub"],
            "Средняя цена за литр, ₽": r["avg_price_per_liter"],
        }
        for r in json_rows
    ]

    if export_rows:
        export_rows.append(
            {
                "Вид топлива": "ИТОГО",
                "Продаж": sum(int(r["sales_count"]) for r in json_rows),
                "Литров": sum(float(r["total_liters"]) for r in json_rows),
                "Сумма, ₽": sum(float(r["total_rub"]) for r in json_rows),
                "Средняя цена за литр, ₽": "",
            }
        )

    if fmt == "csv":
        return _stream_csv(export_rows, "fuel_by_type.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows,
            filename="fuel_by_type.xlsx",
            sheet_name="Топливо по видам",
            columns=["Вид топлива", "Продаж", "Литров", "Сумма, ₽", "Средняя цена за литр, ₽"],
            number_formats={
                "Продаж": "0",
                "Литров": "#,##0.000",
                "Сумма, ₽": "#,##0.00",
                "Средняя цена за литр, ₽": "#,##0.00",
            },
        )
    return json_rows


@router.get("/bonus-movement")
async def report_bonus_movement(
    date_from: date = Query(...),
    date_to: date = Query(...),
    movement_type: Optional[str] = Query(None),
    fmt: Optional[str] = Query(None, alias="format"),
    _: bool = Depends(require_reports_access),
    db: AsyncSession = Depends(get_db),
):
    from sqlalchemy.orm import joinedload
    dt_from, dt_to_exclusive = _report_datetime_bounds(date_from, date_to)
    normalized_movement_type = (movement_type or "").strip().lower() or None
    if normalized_movement_type not in {None, "accrual", "redemption", "expire"}:
        raise HTTPException(status_code=400, detail="Некорректный тип движения бонусов")
    stmt = (
        select(Transaction)
        .options(joinedload(Transaction.client))
        .where(
            Transaction.ts >= dt_from,
            Transaction.ts < dt_to_exclusive,
            ~_cancelled_transaction_clause(),
        )
        .order_by(Transaction.ts.desc())
        .limit(500)
    )
    if normalized_movement_type is not None:
        stmt = stmt.where(Transaction.type == normalized_movement_type)
    result = await db.execute(stmt)
    type_labels = {"accrual": "Начисление", "redemption": "Списание", "expire": "Сгорание"}
    loc_labels = {"fuel": "АЗС", "base": "Магазин"}
    items: list[dict[str, Any]] = []

    # 1) Транзакции бонусов (зарегистрированные клиенты)
    for t in result.unique().scalars().all():
        t_val = t.type.value if hasattr(t.type, 'value') else str(t.type)
        loc_val = t.location.value if t.location and hasattr(t.location, 'value') else str(t.location or '')
        client_name = (t.client.name or '') if t.client else ''
        client_phone = (t.client.phone or '') if t.client else ''

        ts_str = _dt_label(t.ts)
        type_str = type_labels.get(t_val, t_val)
        loc_str = loc_labels.get(loc_val, loc_val)
        amount_bonus = _to_number(t.amount_bonus)
        purchase_amount = _to_number(t.purchase_amount)
        fuel_liters = None if t.fuel_liters is None else _to_number(t.fuel_liters)

        items.append(
            {
                "_ts": t.ts,
                "csv": {
                    "Дата": ts_str,
                    "Телефон": client_phone,
                    "Имя клиента": client_name,
                    "Тип": type_str,
                    "Бонусы": amount_bonus,
                    "Сумма чека, ₽": purchase_amount,
                    "Литры топлива": "" if fuel_liters is None else fuel_liters,
                    "Место": loc_str,
                    "Номер чека": t.check_id or "",
                },
                "xlsx": {
                    "Дата": _excel_datetime(t.ts) if t.ts else ts_str,
                    "Телефон": client_phone,
                    "Имя клиента": client_name,
                    "Тип": type_str,
                    "Бонусы": amount_bonus,
                    "Сумма чека, ₽": purchase_amount,
                    "Литры топлива": fuel_liters,
                    "Место": loc_str,
                    "Номер чека": t.check_id or "",
                },
            }
        )

    # 2) Продажи топлива без клиента (незарегистрированные/не указаны)
    sales_stmt = (
        select(FuelSale)
        .where(
            FuelSale.sale_date >= dt_from,
            FuelSale.sale_date < dt_to_exclusive,
            FuelSale.client_id.is_(None),
            ~_cancelled_fuel_sale_clause(),
        )
        .order_by(FuelSale.sale_date.desc())
        .limit(500)
    )
    sales_result = await db.execute(sales_stmt)
    for s in sales_result.scalars().all():
        ts = s.sale_date
        ts_str = _dt_label(ts)
        purchase_amount = _to_number(s.total_rub)
        fuel_liters = _to_number(s.liters)
        items.append(
            {
                "_ts": ts,
                "csv": {
                    "Дата": ts_str,
                    "Телефон": "",
                    "Имя клиента": "Без клиента",
                    "Тип": "Продажа (без клиента)",
                    "Бонусы": 0.0,
                    "Сумма чека, ₽": purchase_amount,
                    "Литры топлива": fuel_liters,
                    "Место": "АЗС",
                    "Номер чека": "",
                },
                "xlsx": {
                    "Дата": _excel_datetime(ts) if ts else ts_str,
                    "Телефон": "",
                    "Имя клиента": "Без клиента",
                    "Тип": "Продажа (без клиента)",
                    "Бонусы": 0.0,
                    "Сумма чека, ₽": purchase_amount,
                    "Литры топлива": fuel_liters,
                    "Место": "АЗС",
                    "Номер чека": "",
                },
            }
        )

    # Сортировка и лимит (общий)
    items.sort(key=lambda x: (x.get("_ts") is not None, x.get("_ts")), reverse=True)
    items = items[:500]

    export_rows_csv = [i["csv"] for i in items]
    export_rows_xlsx = [i["xlsx"] for i in items]

    export_suffix = {
        None: ("bonus_movement", "Движение бонусов"),
        "accrual": ("bonus_accruals", "Начисления бонусов"),
        "redemption": ("bonus_redemptions", "Списания бонусов"),
        "expire": ("bonus_expire", "Сгорание бонусов"),
    }[normalized_movement_type]

    if fmt == "csv":
        return _stream_csv(export_rows_csv, f"{export_suffix[0]}.csv")
    if fmt == "xlsx":
        return _stream_xlsx(
            rows=export_rows_xlsx,
            filename=f"{export_suffix[0]}.xlsx",
            sheet_name=export_suffix[1],
            columns=[
                "Дата",
                "Телефон",
                "Имя клиента",
                "Тип",
                "Бонусы",
                "Сумма чека, ₽",
                "Литры топлива",
                "Место",
                "Номер чека",
            ],
            number_formats={
                "Дата": "dd.mm.yyyy hh:mm",
                "Бонусы": "#,##0.00",
                "Сумма чека, ₽": "#,##0.00",
                "Литры топлива": "#,##0.000",
            },
        )

    # Для JSON возвращаем прежние ключи (EN) чтобы таблица работала
    json_rows = []
    for r in export_rows_csv:
        json_rows.append(
            {
                "ts": r["Дата"],
                "client_phone": r["Телефон"],
                "client_name": r["Имя клиента"],
                "type": r["Тип"],
                "amount_bonus": r["Бонусы"],
                "purchase_amount": r["Сумма чека, ₽"],
                "fuel_liters": r["Литры топлива"],
                "location": r["Место"],
                "check_id": r["Номер чека"],
            }
        )
    return json_rows


